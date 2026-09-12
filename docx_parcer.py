import os
from collections import defaultdict
from typing import List, Dict, Any
from docx import Document
from docx2python import docx2python
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.text.hyperlink import Hyperlink
from docx.oxml.ns import qn

from interfaces import ParserInterface, ParsedDocument


class DocxParser(ParserInterface):

    def parse(self, file_path: str) -> ParsedDocument:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Файл {file_path} не найден!")

        content = docx2python(file_path)
        footnotes_map = self._parse_notes(content.footnotes)
        endnotes_map = self._parse_notes(content.endnotes)

        document = Document(file_path)
        document_timeline: List[str] = []
        extracted_tables: List[Dict[str, Any]] = []

        for element in document.element.body:

            if element.tag.endswith('tbl'):
                table = Table(element, document)
                table_rows_data = []
                document_timeline.append("\n[НАЧАЛО ТАБЛИЦЫ]\n")

                for row in table.rows:
                    row_cells_text = []
                    for cell in row.cells:
                        cell_text = ""
                        for paragraph in cell.paragraphs:
                            cell_text += self._extract_text_from_paragraph(
                                paragraph, document, footnotes_map, endnotes_map
                            ) + " "
                        row_cells_text.append(cell_text.strip())

                    row_str = " | ".join(row_cells_text)
                    document_timeline.append(row_str)
                    table_rows_data.append(row_cells_text)

                document_timeline.append("\n[КОНЕЦ ТАБЛИЦЫ]\n")
                extracted_tables.append({
                    "data": table_rows_data
                })

            elif element.tag.endswith('paragraph') or element.tag.endswith('}p'):
                paragraph_obj = Paragraph(element, document)
                cleaned_text = self._extract_text_from_paragraph(
                    paragraph_obj, document, footnotes_map, endnotes_map
                )
                if cleaned_text:
                    document_timeline.append(cleaned_text)

        full_text = "\n".join(document_timeline)

        return ParsedDocument(
            filename=os.path.basename(file_path),
            file_type="docx",
            full_text=full_text,
            pages_or_sheets={"main": full_text},
            tables=extracted_tables if extracted_tables else None,
            metadata={
                "paragraphs_count": len(document.paragraphs),
                "tables_count": len(document.tables),
                "footnotes_count": len(footnotes_map),
                "endnotes_count": len(endnotes_map)
            }
        )

    def _parse_notes(self, notes_content: List) -> Dict[int, str]:
        notes_map = defaultdict(str)
        i = 1
        for note in notes_content:
            is_table_open = False
            for blocks in note:
                for paragraph in blocks:
                    if len(paragraph) > 1 and not is_table_open:
                        notes_map[i] += "\n[НАЧАЛО ТАБЛИЦЫ]\n"
                        is_table_open = True
                        notes_map[i] += " | ".join(paragraph)
                    elif is_table_open and len(paragraph) == 1:
                        notes_map[i] += "\n[КОНЕЦ ТАБЛИЦЫ]\n"
                        notes_map[i] += "".join(paragraph) + " "
                        is_table_open = False
                    elif is_table_open and len(paragraph) > 1:
                        notes_map[i] += "\n" + " | ".join(paragraph)
                    else:
                        notes_map[i] += "".join(paragraph) + " "

            notes_map[i] = notes_map[i].strip()
            if is_table_open:
                notes_map[i] += "\n[КОНЕЦ ТАБЛИЦЫ]\n"
            i += 1
        return notes_map

    def _extract_text_from_paragraph(
            self,
            paragraph: Paragraph,
            doc: Document,
            footnotes_map: Dict[int, str],
            endnotes_map: Dict[int, str]
    ) -> str:
        text = ""

        for child in paragraph.iter_inner_content():
            if isinstance(child, Run):
                text += child.text
            elif isinstance(child, Hyperlink):
                visible_text = child.text
                r_id = child._element.get(qn('r:id'))
                url = doc.part.rels[r_id].target_ref if r_id and r_id in doc.part.rels else ""
                text += f'<a href="{url}">{visible_text}</a>'

        w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

        for ref in paragraph._element.findall('.//' + w + 'footnoteReference'):
            f_id = ref.get(qn('w:id'))
            if f_id and f_id.isdigit():
                f_text = footnotes_map.get(int(f_id), "Сноска не найдена")
                text += f" [СНОСКА: {f_text}] "

        for ref in paragraph._element.findall('.//' + w + 'endnoteReference'):
            e_id = ref.get(qn('w:id'))
            if e_id and e_id.isdigit():
                e_text = endnotes_map.get(int(e_id), "Сноска не найдена")
                text += f" [СНОСКА: {e_text}] "

        return text.replace('\n', ' ').strip()


if __name__ == "__main__":
    test_file = "demo.docx"
    if os.path.exists(test_file):
        parser = DocxParser()
        parsed_doc = parser.parse(test_file)

        print("=== ТЕСТ PARSED DOCUMENT ===")
        print(f"Имя файла: {parsed_doc.filename}")
        print(f"Формат: {parsed_doc.file_type}")
        print(f"Метаданные: {parsed_doc.metadata}")
        print("\nФрагмент текста:\n")
        print(parsed_doc.full_text[:300])