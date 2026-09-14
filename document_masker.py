"""
Маскирование конфиденциальных данных в DOCX, XLSX, PDF.

DOCX: run'ы собираются рекурсивно (включая w:hyperlink, w:ins, w:del, w:sdt,
w:smartTag, w:customXml, w:fldSimple). При замене run разбивается так, что
подсветка попадает только на маскируемый фрагмент.

XLSX: формулы не затрагиваются. Маскированные фрагменты внутри значения
ячейки выделяются красным жирным шрифтом через CellRichText, изменённые
ячейки дополнительно помечаются светло-жёлтой заливкой (опционально).

PDF: устойчивый поиск (точный / по пробелам / пословный), дедупликация
прямоугольников, один apply_redactions() на страницу.
"""

import logging
import re
from copy import deepcopy
from functools import lru_cache
from typing import Dict, List, Optional, Tuple

import pymupdf as fitz  # PyMuPDF
import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import PatternFill

logger = logging.getLogger(__name__)


W_R = qn('w:r')
W_T = qn('w:t')
W_RPR = qn('w:rPr')
W_P = qn('w:p')
XML_SPACE = qn('xml:space')


# ---------------------------------------------------------------------------
# Общие утилиты
# ---------------------------------------------------------------------------

def _get_sorted_replacements(replacements: Dict[str, str]) -> Dict[str, str]:
    return dict(sorted(
        replacements.items(),
        key=lambda item: len(item[0]),
        reverse=True,
    ))


@lru_cache(maxsize=4096)
def _build_flexible_regex_cached(target_text: str) -> re.Pattern:
    escaped = re.escape(target_text.strip())
    pattern_str = re.sub(r'\\?\s+', r'\\s+', escaped)
    return re.compile(pattern_str, re.IGNORECASE)


def _build_flexible_regex(target_text: str) -> re.Pattern:
    return _build_flexible_regex_cached(target_text)


def _prepare_patterns(replacements: Dict[str, str]) -> List[Tuple[re.Pattern, str]]:
    prepared: List[Tuple[re.Pattern, str]] = []
    for target, mask in _get_sorted_replacements(replacements).items():
        if not target or not target.strip():
            continue
        prepared.append((_build_flexible_regex(target), mask))
    return prepared


# ===========================================================================
# DOCX
# ===========================================================================

def _collect_runs(paragraph: Paragraph) -> List[Run]:
    """Все w:r внутри абзаца, включая вложенные в гиперссылки и sdt."""
    runs: List[Run] = []

    def _walk(el) -> None:
        for child in el:
            tag = child.tag
            if tag == W_R:
                runs.append(Run(child, paragraph))
            elif tag == W_P:
                continue
            else:
                _walk(child)

    _walk(paragraph._element)
    return runs


def _get_run_rpr(run_el) -> Optional[object]:
    rPr = run_el.find(W_RPR)
    return deepcopy(rPr) if rPr is not None else None


def _run_has_complex_content(run_el) -> bool:
    for child in run_el:
        if child.tag in (W_RPR, W_T):
            continue
        return True
    return False


def _make_run_element(parent, rPr, text: str):
    r_el = parent.makeelement(W_R, {})
    if rPr is not None:
        r_el.append(deepcopy(rPr))
    t_el = parent.makeelement(W_T, {})
    t_el.set(XML_SPACE, 'preserve')
    t_el.text = text
    r_el.append(t_el)
    return r_el


def _collect_non_overlapping_matches(
    full_text: str, patterns: List[Tuple[re.Pattern, str]]
) -> List[Tuple[int, int, str]]:
    raw: List[Tuple[int, int, str]] = []
    for pattern, mask in patterns:
        for match in pattern.finditer(full_text):
            raw.append((match.start(), match.end(), mask))

    if not raw:
        return []

    raw.sort(key=lambda r: (r[0], -(r[1] - r[0])))

    result: List[Tuple[int, int, str]] = []
    last_end = 0
    for start, end, mask in raw:
        if start >= last_end:
            result.append((start, end, mask))
            last_end = end
    return result


def _replace_in_paragraph(
    paragraph: Paragraph, patterns: List[Tuple[re.Pattern, str]]
) -> bool:
    runs = _collect_runs(paragraph)
    if not runs:
        return False

    full_text = "".join(run.text or "" for run in runs)
    if not full_text.strip():
        return False

    matches = _collect_non_overlapping_matches(full_text, patterns)
    if not matches:
        return False

    run_starts: List[int] = []
    pos = 0
    for run in runs:
        run_starts.append(pos)
        pos += len(run.text or "")
    total_len = pos

    run_ops: List[List[Tuple[int, int, str]]] = [[] for _ in runs]

    for m_start, m_end, mask in matches:
        affected: List[Tuple[int, int, int]] = []
        for i, r_start in enumerate(run_starts):
            r_end = run_starts[i + 1] if i + 1 < len(run_starts) else total_len
            if m_end <= r_start or m_start >= r_end:
                continue
            affected.append((i, r_start, r_end))

        if not affected:
            continue

        if any(_run_has_complex_content(runs[i]._element) for i, _, _ in affected):
            logger.warning(
                "Пропуск маскирования '%s': в run служебные элементы",
                full_text[m_start:m_end],
            )
            continue

        for idx, (i, r_start, r_end) in enumerate(affected):
            local_start = max(0, m_start - r_start)
            local_end = min(r_end - r_start, m_end - r_start)
            if idx == 0:
                run_ops[i].append((local_start, local_end, mask))
            else:
                run_ops[i].append((local_start, local_end, ""))

    if not any(run_ops):
        return False

    plans: List[Optional[List[object]]] = [None] * len(runs)
    any_changed = False

    for i, run in enumerate(runs):
        ops = run_ops[i]
        if not ops:
            continue

        run_el = run._element
        run_text = run.text or ""
        rPr = _get_run_rpr(run_el)
        parent = run_el.getparent()

        ops.sort(key=lambda o: o[0])

        segments: List[Tuple[str, bool]] = []
        cursor = 0
        for local_start, local_end, mask in ops:
            if local_start > cursor:
                segments.append((run_text[cursor:local_start], False))
            if mask:
                segments.append((mask, True))
            cursor = local_end
        if cursor < len(run_text):
            segments.append((run_text[cursor:], False))

        new_els: List[object] = []
        for text, is_mask in segments:
            if not text:
                continue
            el = _make_run_element(parent, rPr, text)
            if is_mask:
                wrapper = Run(el, paragraph)
                wrapper.font.highlight_color = WD_COLOR_INDEX.YELLOW
            new_els.append(el)

        plans[i] = new_els
        any_changed = True

    if not any_changed:
        return False

    for i, run in enumerate(runs):
        if plans[i] is None:
            continue
        run_el = run._element
        actual_parent = run_el.getparent()
        if actual_parent is None:
            continue
        for el in plans[i]:
            run_el.addprevious(el)
        actual_parent.remove(run_el)

    return True


def _process_docx_paragraph(
    paragraph: Paragraph, patterns: List[Tuple[re.Pattern, str]]
) -> None:
    try:
        _replace_in_paragraph(paragraph, patterns)
    except Exception as exc:
        logger.warning("Не удалось обработать абзац DOCX: %s", exc)


def mask_docx(input_path: str, output_path: str, replacements: Dict[str, str]) -> None:
    doc = docx.Document(input_path)
    patterns = _prepare_patterns(replacements)

    for paragraph in doc.paragraphs:
        _process_docx_paragraph(paragraph, patterns)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _process_docx_paragraph(paragraph, patterns)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            _process_docx_paragraph(paragraph, patterns)
        for paragraph in section.footer.paragraphs:
            _process_docx_paragraph(paragraph, patterns)

    doc.save(output_path)


# ===========================================================================
# XLSX  (пофрагментная разметка через CellRichText)
# ===========================================================================

# Стили для маскированного и обычного текста внутри ячейки.
# Внимание: openpyxl не даёт задать фоновую подсветку отдельного фрагмента,
# поэтому маска выделяется насыщенным красным жирным шрифтом, а изменённая
# ячейка целиком получает светло-жёлтую заливку.
_MASK_FONT = InlineFont(
    rFont="Calibri",
    sz=11,
    b=True,
    color="FFC00000",       # тёмно-красный
)
_PLAIN_FONT = InlineFont(
    rFont="Calibri",
    sz=11,
    color="FF000000",       # чёрный
)

_MODIFIED_FILL = PatternFill(
    start_color="FFFFF2CC",
    end_color="FFFFF2CC",
    fill_type="solid",
)


def _cell_is_formula(cell) -> bool:
    """True, если ячейка содержит формулу — такие не трогаем."""
    if cell.data_type == 'f':
        return True
    if isinstance(cell.value, str) and cell.value.startswith('='):
        return True
    return False


def _find_non_overlapping_matches_xlsx(
    text: str,
    patterns: List[Tuple[re.Pattern, str]],
) -> List[Tuple[int, int, str]]:
    """Непересекающиеся (start, end, mask) по возрастанию позиции.

    При пересечении приоритет у самой длинной замены.
    """
    raw: List[Tuple[int, int, str]] = []
    for pattern, mask in patterns:
        for m in pattern.finditer(text):
            raw.append((m.start(), m.end(), mask))

    if not raw:
        return []

    raw.sort(key=lambda r: (r[0], -(r[1] - r[0])))

    result: List[Tuple[int, int, str]] = []
    last_end = 0
    for start, end, mask in raw:
        if start >= last_end:
            result.append((start, end, mask))
            last_end = end
    return result


def _build_rich_value(
    original_text: str,
    matches: List[Tuple[int, int, str]],
) -> CellRichText:
    """CellRichText: маска — красным жирным, остальной текст — обычным."""
    blocks: List[TextBlock] = []
    cursor = 0

    for start, end, mask in matches:
        if start > cursor:
            blocks.append(TextBlock(_PLAIN_FONT, original_text[cursor:start]))
        blocks.append(TextBlock(_MASK_FONT, mask))
        cursor = end

    if cursor < len(original_text):
        blocks.append(TextBlock(_PLAIN_FONT, original_text[cursor:]))

    return CellRichText(blocks)


def mask_xlsx(
    input_path: str,
    output_path: str,
    replacements: Dict[str, str],
    highlight_cells: bool = True,
) -> None:
    """Маскирует конфиденциальные данные в XLSX.

    Args:
        input_path: Путь к исходному .xlsx.
        output_path: Путь для сохранения обезличенного .xlsx.
        replacements: Словарь {оригинальный_текст: маска}.
        highlight_cells: Помечать ли изменённые ячейки светло-жёлтой
            заливкой. Маскированные фрагменты внутри текста всегда
            выделяются красным жирным шрифтом.
    """
    wb = openpyxl.load_workbook(input_path)
    patterns = _prepare_patterns(replacements)

    if not patterns:
        wb.save(output_path)
        logger.info("XLSX: список замен пуст, файл сохранён без изменений.")
        return

    total_replacements = 0
    total_cells = 0

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if _cell_is_formula(cell):
                    continue

                cell_text = str(cell.value)
                matches = _find_non_overlapping_matches_xlsx(cell_text, patterns)
                if not matches:
                    continue

                cell.value = _build_rich_value(cell_text, matches)

                if highlight_cells:
                    cell.fill = _MODIFIED_FILL

                total_replacements += len(matches)
                total_cells += 1

    wb.save(output_path)
    logger.info(
        "XLSX: заменено %d фрагментов в %d ячейках. Файл: %s",
        total_replacements, total_cells, output_path,
    )


# ===========================================================================
# PDF
# ===========================================================================

def _search_phrase_fuzzy(page: "fitz.Page", phrase: str) -> List["fitz.Rect"]:
    """Ищет фразу на странице PDF, устойчиво к различиям в пробелах.

    1. Точный поиск.
    2. С нормализованными пробелами.
    3. Пословный поиск в пределах одной строки.
    """
    phrase = (phrase or "").strip()
    if not phrase:
        return []

    rects = page.search_for(phrase)
    if rects:
        return rects

    normalized = re.sub(r'\s+', ' ', phrase)
    if normalized != phrase:
        rects = page.search_for(normalized)
        if rects:
            return rects

    words = phrase.split()
    if len(words) < 2:
        return []

    try:
        page_words = page.get_text("words")
    except Exception:
        return []

    def _norm(w: str) -> str:
        return re.sub(r'\W+', '', w).lower()

    target = [_norm(w) for w in words]
    result: List["fitz.Rect"] = []

    for i in range(len(page_words) - len(words) + 1):
        window = page_words[i:i + len(words)]
        if len({(w[5], w[6]) for w in window}) != 1:
            continue
        if [_norm(w[4]) for w in window] != target:
            continue

        x0 = min(w[0] for w in window)
        y0 = min(w[1] for w in window)
        x1 = max(w[2] for w in window)
        y1 = max(w[3] for w in window)
        result.append(fitz.Rect(x0, y0, x1, y1))

    return result


def _collect_rects_with_masks(
    page: "fitz.Page", sorted_replacements: Dict[str, str]
) -> List[Tuple["fitz.Rect", str]]:
    seen: set = set()
    result: List[Tuple["fitz.Rect", str]] = []

    for target, mask in sorted_replacements.items():
        if not target or not target.strip():
            continue

        for rect in _search_phrase_fuzzy(page, target):
            key = (
                round(rect.x0, 1),
                round(rect.y0, 1),
                round(rect.x1, 1),
                round(rect.y1, 1),
            )
            if key in seen:
                continue
            seen.add(key)
            result.append((rect, mask))

    return result


def mask_pdf(
    input_path: str,
    output_path: str,
    replacements: Dict[str, str],
    add_text_overlay: bool = True,
) -> None:
    doc = fitz.open(input_path)
    sorted_replacements = _get_sorted_replacements(replacements)

    for page_num, page in enumerate(doc, 1):
        rects_with_masks = _collect_rects_with_masks(page, sorted_replacements)
        logger.info(
            "Страница %d: найдено %d областей для маскирования",
            page_num, len(rects_with_masks),
        )

        for rect, mask in rects_with_masks:
            if add_text_overlay:
                page.add_redact_annot(
                    rect,
                    text=mask,
                    fill=(1, 1, 0),
                    text_color=(0, 0, 0),
                )
            else:
                page.add_redact_annot(rect, fill=(1, 1, 0))

        page.apply_redactions()

    doc.save(output_path)