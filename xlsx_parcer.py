import os
from typing import List, Dict, Any
import openpyxl

from interfaces import ParserInterface, ParsedDocument


class XlsxParser(ParserInterface):

    def parse(self, file_path: str) -> ParsedDocument:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Файл {file_path} не найден!")

        wb = openpyxl.load_workbook(file_path, data_only=True)

        pages_or_sheets: Dict[str, str] = {}
        full_text_list: List[str] = []
        extracted_tables: List[Dict[str, Any]] = []

        for sheet in wb.worksheets:
            sheet_lines: List[str] = []
            sheet_matrix: List[List[str]] = []

            for row in sheet.iter_rows():
                row_vals = []
                for cell in row:
                    val = str(cell.value) if cell.value is not None else ""
                    row_vals.append(val.strip())

                if any(row_vals):
                    non_empty_vals = [v for v in row_vals if v]
                    sheet_lines.append(" | ".join(non_empty_vals))
                    sheet_matrix.append(row_vals)

            sheet_text = "\n".join(sheet_lines)
            pages_or_sheets[sheet.title] = sheet_text
            full_text_list.append(f"=== Лист: {sheet.title} ===\n" + sheet_text)

            extracted_tables.append({
                "sheet_name": sheet.title,
                "data": sheet_matrix
            })

        full_text = "\n\n".join(full_text_list)

        return ParsedDocument(
            filename=os.path.basename(file_path),
            file_type="xlsx",
            full_text=full_text,
            pages_or_sheets=pages_or_sheets,
            tables=extracted_tables if extracted_tables else None,
            metadata={
                "sheet_count": len(wb.worksheets),
                "sheet_names": wb.sheetnames
            }
        )


if __name__ == "__main__":
    test_xlsx = "sample_table.xlsx"
    if os.path.exists(test_xlsx):
        parser = XlsxParser()
        parsed_doc = parser.parse(test_xlsx)

        print("=== РЕЗУЛЬТАТ ПАРСИНГА XLSX ===")
        print(f"Имя файла: {parsed_doc.filename}")
        print(f"Количество листов: {parsed_doc.metadata.get('sheet_count')}")
        print(f"Имена листов: {parsed_doc.metadata.get('sheet_names')}")
        print("\nФрагмент текста:\n")
        print(parsed_doc.full_text[:300])